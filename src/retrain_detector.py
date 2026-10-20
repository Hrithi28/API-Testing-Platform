"""
Retrain Module 4's supervised detector on VERIFIED labels and evaluate whether it
generalises to a service it has never seen.

Compared on the same verified ground truth (label_findings.py):
  A  original rule-based finding decision   (is the response flagged by M4's rules?)
  B  original weak-label Random Forest      (trained on heuristic_label, tested cross-target)
  C  verified-label Random Forest           (trained on verified labels, tested cross-target)
  C2 verified-label RF on a PORTABLE feature set (no raw status code / latency / size,
     which are service-specific). Chosen after observing C fail on httpbin -- disclosed
     as a post-hoc design decision.
  D  the labelling rule itself applied as a classifier (reachable only because the
     labels are a function of observable signals) -- a reference, not a competitor
  C' verified-label RF, pooled stratified 5-fold CV   (optimistic: near-duplicate rows)

Cross-target protocol: train on one service, test on the other, both directions.
That is the generalisation test the original circular evaluation could not provide.

Reviewer overrides: if artifacts_real/label_review.xlsx exists, any value (1/0) in
its 'reviewer_override' column replaces the rubric label for that response.
"""
from __future__ import annotations
import json
import os
import sys
from typing import Any, Dict, List, Tuple

import numpy as np
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import precision_recall_fscore_support
from sklearn.model_selection import StratifiedKFold

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from m3_executor import ExecutionResult                      # noqa: E402
from m4_anomaly_detector import featurize, heuristic_label    # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ART = os.path.join(ROOT, "artifacts_real")
TARGETS = ["httpbin", "chroma"]
CATS = ["positive", "negative", "boundary", "security"]
RF_FIELDS = set(ExecutionResult.__dataclass_fields__)
FEATS = ["status_code", "is_server_error", "is_client_error", "is_success", "latency_ms", "log_latency",
         "response_size", "n_schema_violations", "status_unexpected", "is_security_probe",
         "echoes_probe_payload", "is_json", "cat_positive", "cat_negative", "cat_boundary",
         "cat_security", "is_write", "is_html"]
PORTABLE = [FEATS.index(n) for n in FEATS
            if n not in ("status_code", "latency_ms", "log_latency", "response_size")]


def extra_features(r: ExecutionResult, method: str) -> List[float]:
    return [1.0 if r.category == c else 0.0 for c in CATS] + [
        1.0 if method in ("POST", "PUT", "PATCH", "DELETE") else 0.0,
        1.0 if "html" in (r.content_type or "").lower() else 0.0,
    ]


def load_overrides() -> Dict[Tuple[str, str], int]:
    path = os.path.join(ART, "label_review.xlsx")
    out: Dict[Tuple[str, str], int] = {}
    if not os.path.exists(path):
        return out
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Findings"]
    hdr = [c.value for c in ws[1]]
    ti, ei, oi = hdr.index("target"), hdr.index("test_id"), hdr.index("reviewer_override")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[oi] in (0, 1, "0", "1"):
            out[(row[ti], row[ei])] = int(row[oi])
    return out


def load() -> Dict[str, Dict[str, Any]]:
    labels = json.load(open(os.path.join(ART, "labels.json")))
    over = load_overrides()
    data: Dict[str, Dict[str, Any]] = {}
    for t in TARGETS:
        rep = json.load(open(os.path.join(ART, t, "run_report.json")))
        meth = {c["test_id"]: c["method"] for c in rep["test_cases"]}
        flagged = {f["test_id"] for f in rep["with_oracle"]["findings"]}
        lab = {x["test_id"]: x["rubric_label"] for x in labels if x["target"] == t}
        X, Xb, y, yw, rule_flag, ids, dflag = [], [], [], [], [], [], []
        for rd in rep["results"]:
            tid = rd["test_id"]
            v = over.get((t, tid), lab.get(tid))
            if v is None:
                continue                         # excluded: non-deterministic / awaiting review
            r = ExecutionResult(**{k: rd[k] for k in rd if k in RF_FIELDS})
            base = featurize(r)
            X.append(base + extra_features(r, meth[tid]))
            Xb.append(base)
            y.append(int(v)); yw.append(heuristic_label(r))
            rule_flag.append(1 if tid in flagged else 0); ids.append(tid)
            hard = [v for v in r.schema_violations if not v.startswith("semantic:")]
            dflag.append(1 if ((500 <= r.status_code < 600 and r.status_code != 501)
                               or (hard and 200 <= r.status_code < 300)) else 0)
        data[t] = {"X": np.array(X), "Xb": np.array(Xb), "y": np.array(y), "yw": np.array(yw),
                   "flag": np.array(rule_flag), "dflag": np.array(dflag), "ids": ids}
    return data


def prf(y, p) -> Dict[str, float]:
    pr, rc, f1, _ = precision_recall_fscore_support(y, p, average="binary", zero_division=0)
    tp = int(((y == 1) & (p == 1)).sum()); fp = int(((y == 0) & (p == 1)).sum())
    fn = int(((y == 1) & (p == 0)).sum()); tn = int(((y == 0) & (p == 0)).sum())
    return {"precision": round(float(pr), 4), "recall": round(float(rc), 4), "f1": round(float(f1), 4),
            "tp": tp, "fp": fp, "fn": fn, "tn": tn, "n": int(len(y)), "positives": int(y.sum())}


def rf(seed: int = 42) -> RandomForestClassifier:
    return RandomForestClassifier(n_estimators=300, random_state=seed,
                                  class_weight="balanced", min_samples_leaf=1)


def main() -> None:
    D = load()
    out: Dict[str, Any] = {"label_counts": {t: {"n": int(len(D[t]["y"])), "defects": int(D[t]["y"].sum())}
                                              for t in TARGETS},
                           "cross_target": {}, "pooled_cv": {}}
    for tr, te in (("httpbin", "chroma"), ("chroma", "httpbin")):
        a, b = D[tr], D[te]
        res = {}
        res["A_rule_flag"] = prf(b["y"], b["flag"])
        mB = rf().fit(a["Xb"], a["yw"])                       # original: weak labels, 12 features
        res["B_weak_label_rf"] = prf(b["y"], mB.predict(b["Xb"]))
        mC = rf().fit(a["X"], a["y"])                         # verified labels, 18 features
        res["C_verified_rf"] = prf(b["y"], mC.predict(b["X"]))
        mC2 = rf().fit(a["X"][:, PORTABLE], a["y"])
        res["C2_portable_rf"] = prf(b["y"], mC2.predict(b["X"][:, PORTABLE]))
        res["D_rubric_rule"] = prf(b["y"], b["dflag"])
        res["C_top_features"] = sorted(
            [(n, round(float(w), 3)) for n, w in zip(FEATS, mC.feature_importances_)],
            key=lambda kv: -kv[1])[:5]
        out["cross_target"][f"train_{tr}__test_{te}"] = res
    # pooled CV (optimistic)
    X = np.vstack([D[t]["X"] for t in TARGETS]); y = np.concatenate([D[t]["y"] for t in TARGETS])
    runs = []
    for seed in range(5):
        P = np.zeros_like(y)
        for tri, tei in StratifiedKFold(5, shuffle=True, random_state=seed).split(X, y):
            P[tei] = rf(seed).fit(X[tri], y[tri]).predict(X[tei])
        runs.append(prf(y, P))
    out["pooled_cv"] = {"mean_precision": round(float(np.mean([r["precision"] for r in runs])), 4),
                        "mean_recall": round(float(np.mean([r["recall"] for r in runs])), 4),
                        "mean_f1": round(float(np.mean([r["f1"] for r in runs])), 4),
                        "sd_f1": round(float(np.std([r["f1"] for r in runs])), 4), "runs": 5}
    json.dump(out, open(os.path.join(ART, "retrain_results.json"), "w"), indent=1)

    print("labelled responses:", out["label_counts"])
    for k, v in out["cross_target"].items():
        print(f"\n{k}")
        for m in ("A_rule_flag", "B_weak_label_rf", "C_verified_rf", "C2_portable_rf", "D_rubric_rule"):
            r = v[m]
            print(f"  {m:18s} P={r['precision']:.3f} R={r['recall']:.3f} F1={r['f1']:.3f}  "
                  f"(TP={r['tp']} FP={r['fp']} FN={r['fn']} TN={r['tn']})")
        print("  top features:", v["C_top_features"])
    print("\npooled 5-fold CV (optimistic):", out["pooled_cv"])


if __name__ == "__main__":
    main()
