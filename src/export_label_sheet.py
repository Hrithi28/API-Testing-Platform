"""Export every raised finding with its rubric verdict to an Excel workbook for human review.
Type 1 (real defect) or 0 (not a defect) in 'reviewer_override' to overrule a verdict; retrain_detector.py honours it."""
import json, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "artifacts_real")
rows = [x for x in json.load(open(f"{A}/labels.json")) if x["is_finding"]]
rows.sort(key=lambda x: (x["target"], {"R1": 0, "R2": 0, "R1x": 1, "REVIEW": 1}.get(x["rule"], 2), x["endpoint"], x["test_id"]))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Findings"
cols = ["target", "test_id", "endpoint", "category", "intent", "status", "declared_codes", "response_preview",
        "finding_class", "rule", "rubric_label", "rationale", "repro_statuses", "reviewer_override", "reviewer_note"]
ws.append(cols)
for x in rows:
    clean = lambda v: re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "?", v) if isinstance(v, str) else v
    ws.append([clean(", ".join(map(str, x[c])) if isinstance(x[c], list) else x[c]) for c in cols[:-2]] + [None, None])
hdr = PatternFill("solid", fgColor="335295")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(vertical="center", wrap_text=True)
widths = {"target": 9, "test_id": 26, "endpoint": 44, "category": 10, "intent": 40, "status": 7, "declared_codes": 14,
          "response_preview": 50, "finding_class": 13, "rule": 8, "rubric_label": 8, "rationale": 60,
          "repro_statuses": 12, "reviewer_override": 12, "reviewer_note": 30}
for i, c in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths[c]
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
yellow = PatternFill("solid", fgColor="FFF2CC")
for r in range(2, ws.max_row + 1):
    ws.cell(r, cols.index("reviewer_override") + 1).fill = yellow

g = wb.create_sheet("How to review")
for line in [
    "Each row is a finding the platform raised. 'rubric_label' is 1 for an actionable defect, 0 otherwise.",
    "To disagree with a verdict, type 1 or 0 in the yellow 'reviewer_override' column. Leave blank to accept.",
    "",
    "Rules:",
    "R1   5xx (not 501) that recurs 3/3 times when re-issued alone, in sequence  -> defect",
    "R1x  5xx that did not recur sequentially (concurrency-dependent)            -> excluded from training/evaluation",
    "R2   2xx response that breaks the operation's declared response schema      -> defect",
    "R3   status code not declared for the operation                             -> documentation gap, not counted as a defect",
    "R0   everything else: 501 unimplemented, auth disabled by default, echo-by-design, declared status -> not a defect",
    "",
    "Suggested spot check: filter rule = R1 or R2 and open the request URL against the running service;",
    "then filter rule = R0 and sample ~30 rows to confirm they are really expected behaviour.",
    "After editing, re-run:  python3 src/retrain_detector.py   (it reads artifacts_real/label_review.xlsx)",
]:
    g.append([line])
g.column_dimensions["A"].width = 120
wb.save(f"{A}/label_review.xlsx"); print("rows", len(rows))
